"""Tests for text extraction.

These use in-memory bytes only. No fixture files on disk, so the suite runs
in CI with nothing checked out but the source.
"""

import io
import pytest
from docx import Document

from backend.file_parser import extract_text_from_file


def test_plain_text():
    assert extract_text_from_file(b"Hello world", "resume.txt") == "Hello world"


def test_markdown_and_csv_treated_as_text():
    assert extract_text_from_file(b"# CV", "resume.md") == "# CV"
    assert extract_text_from_file(b"a,b", "data.csv") == "a,b"


def test_extension_is_case_insensitive():
    assert extract_text_from_file(b"Hello", "RESUME.TXT") == "Hello"


def test_invalid_utf8_does_not_crash():
    """errors='replace' is deliberate: a CV with one bad byte should still be
    scored rather than failing the whole batch."""
    out = extract_text_from_file(b"caf\xff", "resume.txt")
    assert out.startswith("caf")


def test_unsupported_extension_raises():
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text_from_file(b"data", "resume.pages")


def test_docx_paragraphs_extracted():
    doc = Document()
    doc.add_paragraph("Ivy Chen")
    doc.add_paragraph("")           # blank paragraphs are skipped
    doc.add_paragraph("FastAPI, Python")
    buf = io.BytesIO()
    doc.save(buf)

    out = extract_text_from_file(buf.getvalue(), "cv.docx")
    assert "Ivy Chen" in out
    assert "FastAPI, Python" in out
    assert "\n\n" not in out        # no gap left by the blank paragraph


def test_docx_tables_extracted():
    """Table cells are joined with ' | '. Plenty of CVs put the skills matrix
    in a table, and dropping it silently loses the most scoreable content."""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "5 years"
    buf = io.BytesIO()
    doc.save(buf)

    out = extract_text_from_file(buf.getvalue(), "cv.docx")
    assert "Python | 5 years" in out


def test_corrupt_pdf_raises_value_error():
    with pytest.raises(ValueError, match="Failed to parse PDF"):
        extract_text_from_file(b"not a pdf at all", "resume.pdf")
